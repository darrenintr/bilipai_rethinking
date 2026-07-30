from pptx import Presentation
p = Presentation('slides/output/paladala-intro-2026.pptx')
W = p.slide_width / 914400.0
H = p.slide_height / 914400.0
print(f'Canvas: {W:.3f} x {H:.3f}')
overflow_count = 0
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        try:
            x = sh.left / 914400.0
            y = sh.top / 914400.0
            w = sh.width / 914400.0
            h = sh.height / 914400.0
        except Exception:
            continue
        if x + w > 10.001 or y + h > 5.626:
            print(f'  slide {i}: shape {sh.shape_type} at ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f} -- overflow!')
            overflow_count += 1
print(f'Overflow shapes: {overflow_count}')
